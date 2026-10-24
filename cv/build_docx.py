#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
توليد نسخة Word احترافية نظيف (RTL) من السيرة الذاتية — القالب «الاحترافي النظيف».
عمود واحد، مساحات بيضاء، خطوط نظيفة، عناوين ملوّنة — بدون بطاقات أو تكدّس.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import os

INK    = RGBColor(0x14, 0x20, 0x30)
BODY   = RGBColor(0x2C, 0x3C, 0x52)
MUTED  = RGBColor(0x68, 0x76, 0x88)
ACCENT = RGBColor(0x18, 0x46, 0x6E)
ACC2   = RGBColor(0x1E, 0x54, 0x80)
FONT   = "Cairo"

doc = Document()

# ---------- إعداد الصفحة ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.0)
sec.left_margin, sec.right_margin = Cm(1.4), Cm(1.4)
# اتجاه الصفحة من اليمين إلى اليسار
sPr = sec._sectPr
bidi = OxmlElement('w:bidi'); sPr.append(bidi)

style = doc.styles['Normal']
style.font.name = FONT
style.font.size = Pt(10)
style.font.color.rgb = BODY
rpr = style.element.get_or_add_rPr()
rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = OxmlElement('w:rFonts'); rpr.append(rf)
for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
    rf.set(qn(a), FONT)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = 1.15


def _run(p, text, size=10, bold=False, color=BODY, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    f = rPr.find(qn('w:rFonts'))
    if f is None:
        f = OxmlElement('w:rFonts'); rPr.append(f)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        f.set(qn(a), FONT)
    rtl = OxmlElement('w:rtl'); rPr.append(rtl)
    if bold:
        b = OxmlElement('w:bCs'); rPr.append(b)
    return r


def add_link(p, text, url, size=9, bold=False, color=ACCENT):
    """إضافة ارتباط تشعبي (يُفتح بالنقر) داخل فقرة RTL."""
    part = p.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        f.set(qn(a), FONT)
    rPr.append(f)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '%02X%02X%02X' % (color[0], color[1], color[2])); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(size * 2))); rPr.append(szcs)
    if bold:
        rPr.append(OxmlElement('w:b')); rPr.append(OxmlElement('w:bCs'))
    rPr.append(OxmlElement('w:rtl'))
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


def para(size=10, bold=False, color=BODY, before=0, after=0, align=None, ls=1.15):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = ls
    if align is not None:
        p.alignment = align
    return p


def border_bottom(p, color="D9E0EA", size=6):
    """خط سفلي دقيق لفقرة."""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def shade(p, color="18466E"):
    """تظليل خلفية فقرة (للعناوين)."""
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    pPr.append(shd)


def no_borders(table):
    """إزالة حدود الجدول (للقوائم النظيفة)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


def light_bottom(table):
    """حدود سفلية فاتحة فقط داخل الجدول."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    insideH = OxmlElement('w:insideH')
    insideH.set(qn('w:val'), 'single')
    insideH.set(qn('w:sz'), '4')
    insideH.set(qn('w:space'), '0')
    insideH.set(qn('w:color'), 'E2E8F1')
    borders.append(insideH)
    tblPr.append(borders)


def dots(n):
    on = "●" * max(0, min(5, int(n)))
    off = "○" * (5 - len(on))
    return on + off


# ===================== الترويسة =====================
p = para(before=0)
_run(p, "محمد البخيتي", size=20, bold=True, color=INK)

p = para(before=2, after=2)
_run(p, "موظف مكتب خدمات عامة — دعم إداري وتنظيم مكتبي", size=11, bold=True, color=ACCENT)

p = para(before=2)
_run(p, "سنتان خبرة في إعداد المستندات وتلخيص المعاملات وخدمة المراجعين  •  إجادة تامة للحاسب والأوفيس  •  سرعة تعلّم استثنائية",
     size=8.4, color=MUTED)

p = para(before=5, after=1)
add_link(p, "الجوال: +966 56 674 3293", "tel:+966566743293", size=8.3, bold=False, color=ACCENT)
_run(p, "   |   ", size=8.3, color=MUTED)
add_link(p, "واتساب", "https://wa.me/966566743293", size=8.3, color=ACCENT)
_run(p, "   |   البريد: ", size=8.3, color=MUTED)
add_link(p, "Mohammedalbkhty@gmail.com", "mailto:Mohammedalbkhty@gmail.com", size=8.3, color=ACCENT)

p = para(before=1)
add_link(p, "لينكدإن", "https://www.linkedin.com/in/mohammed-albukhaiti-215a41302/", size=8.2, color=ACCENT)
_run(p, "   |   الموقع الشخصي: ", size=8.2, color=MUTED)
add_link(p, "engmohammeda.github.io/portfolio", "https://engmohammeda.github.io/portfolio/", size=8.2, color=ACCENT)
_run(p, "   |   الرياض — غرب الرياض", size=8.2, color=MUTED)

p = para(before=1)
_run(p, "العمر: 24 سنة   |   مقيم في الرياض — متاح للعمل فورًا", size=8.2, color=MUTED)

# شريط دقيق تحت الترويسة
p = para(before=4)
border_bottom(p, color="18466E", size=12)
_run(p, "", size=1)


# ===================== الأقسام =====================
def section_heading(title, en=""):
    p = para(before=10, after=4)
    _run(p, title, size=11, bold=True, color=ACC2)
    if en:
        _run(p, "     " + en, size=6.5, color=RGBColor(0x9A, 0xAE, 0xC0))
    border_bottom(p, color="E2E8F1", size=4)


# ---------------- الملف المهني ----------------
section_heading("الملف المهني", "PROFILE")
summary = (
    "موظف مكتب خدمات عامة منظم وخدوم، خبرة سنتين في مكتب أعمال ومكتبة: "
    "أُعدّ المستندات والخطابات وأنسّقها وأدقّقها، وأُلخّص المعاملات في ملخصات واضحة، "
    "وأُدخل البيانات وأؤرشفها، وأستقبل المراجعين وأتابع طلباتهم حتى الإنجاز. "
    "أُجيد الحاسب وبرامج الأوفيس (Word وExcel وPowerPoint والبريد) إجادة تامة، "
    "وأستخدم أدوات الذكاء الاصطناعي في تلخيص المعاملات وصياغة الخطابات وتنظيم العمل. "
    "أتمتّع بحسن التعامل والأخلاق والالتزام، وأتعلّم أي نظام جديد بسرعة — "
    "وأبحث عن استقرار مهني طويل في غرب الرياض."
)
p = para(after=0, ls=1.3)
_run(p, summary, size=9.3)

# ---------------- المهارات ----------------
section_heading("المهارات الأساسية", "CORE SKILLS")

skill_groups = [
    ("الحاسب وبرامج الأوفيس", [
        ("Word — صياغة وتنسيق المستندات والخطابات", 5),
        ("Excel — جداول، معادلات، تقارير", 4),
        ("PowerPoint — عروض تقديمية", 4),
        ("Outlook و Gmail — بريد ومراسلات", 4),
        ("Google Docs / Sheets / Drive — عمل سحابي", 4),
    ]),
    ("الأعمال المكتبية والسكرتارية", [
        ("صياغة الخطابات والمستندات الرسمية", 5),
        ("تلخيص المعاملات وتجهيز الملخصات", 5),
        ("الأرشفة الورقية والإلكترونية", 4),
        ("إدخال البيانات وتدقيقها", 4),
    ]),
    ("خدمة العملاء والتعامل مع الجمهور", [
        ("استقبال المراجعين والرد على الاستفسارات", 5),
        ("الاستماع الفعّال وحل الاعتراضات", 4),
        ("متابعة الطلبات حتى الإنجاز", 5),
        ("السرية والأمانة والالتزام", 5),
    ]),
    ("المنصات الحكومية والأنظمة", [
        ("سرعة تعلّم أي منصة أو نظام جديد", 5),
        ("أبشر، قوى، ناجز، اعتماد — جاهز للعمل عليها", 4),
        ("إدخال الطلبات ومتابعة الحالة وتحديثها", 4),
        ("رفع المستندات إلكترونيًا والتأكد من اكتمالها", 4),
    ]),
    ("أدوات الذكاء الاصطناعي", [
        ("ChatGPT و Gemini و Copilot", 4),
        ("تلخيص المستندات الطويلة واستخراج النقاط", 4),
        ("صياغة الخطابات والمراسلات", 4),
        ("أتمتة المهام المتكررة وتنظيم البيانات", 3),
    ]),
]

for gtitle, items in skill_groups:
    p = para(before=5, after=2)
    _run(p, gtitle, size=8.5, bold=True, color=ACCENT)
    rows = [items[i:i + 2] for i in range(0, len(items), 2)]
    for pair in rows:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit = False
        no_borders(tbl)
        widths = [Cm(8.3), Cm(8.3)]
        for col_idx, it in enumerate(pair):
            cell = tbl.cell(0, col_idx)
            cell.width = widths[col_idx]
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.0
            _run(cp, dots(it[1]) + "   ", size=8.8, color=ACCENT)
            _run(cp, it[0], size=8.8, color=BODY)
            if len(pair) == 1:
                break

# ---------------- الخبرة ----------------
section_heading("الخبرة العملية", "EXPERIENCE")


def job(title, org, meta, items, first=False):
    p = para(before=(0 if first else 6), after=1)
    _run(p, title, size=10, bold=True, color=INK)
    _run(p, "     " + meta, size=8.2, color=MUTED)
    p = para(before=0, after=1)
    _run(p, "— " + org, size=9.2, bold=True, color=ACCENT)
    for it in items:
        p = para(before=0, after=0.5, ls=1.18)
        p.paragraph_format.left_indent = Cm(0.4)
        _run(p, "•  ", size=8.9, bold=True, color=ACCENT)
        _run(p, it, size=8.9)


job("موظف إداري ومكتبي — خدمات المستندات", "مكتبة ومركز خدمات",
    "الرياض • 2023 – 2025 (سنتان)", [
    "إعداد وتنسيق المستندات والخطابات الرسمية بالعربية والإنجليزية ومراجعتها قبل الطباعة.",
    "تلخيص المعاملات وتجهيز ملخصات واضحة عنها، ما رفع سرعة المراجعة والإنجاز.",
    "إدخال البيانات وتنظيم السجلات على Excel والأرشفة الورقية والإلكترونية.",
    "استقبال المراجعين والرد على استفساراتهم ومتابعة طلباتهم حتى إتمامها بصورة مرضية.",
    "توظيف أدوات الذكاء الاصطناعي في الصياغة والتلخيص وتجهيز قوالب للمستندات المتكررة.",
], first=True)

job("مسؤول مبيعات وكاشير", "متجر تجزئة", "الرياض • 2025", [
    "تعامل يومي مباشر مع العملاء: استقبال، شرح، معالجة اعتراضات، وإتمام الطلب بلباقة.",
    "إدخال البيانات وإعداد تقارير المبيعات ومطابقة الأرقام بدقة على Excel.",
])

# ---------------- المؤهل العلمي ----------------
section_heading("المؤهل العلمي", "EDUCATION")
p = para(before=2, after=1)
_run(p, "دبلوم لمدة سنتين", size=9.8, bold=True, color=INK)
_run(p, "     2026", size=8.3, color=MUTED)
p = para(before=0)
_run(p, "التخصص: ........ — اسم الجهة التعليمية", size=8.6, bold=True, color=ACCENT)
p = para(before=0)
_run(p, "خلفية تعزز الدقة والانضباط وتنظيم البيانات والتعامل مع الأرقام.", size=8.4, color=MUTED)

# ---------------- الدورات ----------------
section_heading("الدورات والشهادات", "CERTIFICATES")
p = para(before=2)
for t in ["الحاسب الآلي / ICDL", "Excel للمبتدئين", "Word — المستندات الاحترافية",
          "خدمة العملاء", "أدوات الذكاء الاصطناعي"]:
    _run(p, "•  ", size=8.6, bold=True, color=ACCENT)
    _run(p, t, size=8.6, color=BODY)
    _run(p, "   ", size=8.6)

# ---------------- اللغات والمهارات الشخصية ----------------
section_heading("اللغات والمهارات الشخصية", "LANGUAGES & SOFT SKILLS")
rows = [
    ("العربية", "اللغة الأم — إجادة تامة تحدثًا وكتابة"),
    ("الإنجليزية", "جيد — قراءة وكتابة المراسلات والمستندات"),
    ("سرعة التعلّم", "أتعلّم الأنظمة الجديدة بسرعة وأطوّر نفسي باستمرار"),
    ("الشخصية", "أمانة، التزام، لباقة، تنظيم، روح الفريق"),
    ("تحت الضغط", "أداء ثابت مع كثرة المراجعين وتعدد المهام"),
    ("المظهر والحضور", "مظهر مهني لائق واحترام المواعيد"),
]
i = 0
while i < len(rows):
    pair = rows[i:i + 2]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    light_bottom(tbl)
    widths = [Cm(8.3), Cm(8.3)]
    for col_idx, (k, v) in enumerate(pair):
        cell = tbl.cell(0, col_idx)
        cell.width = widths[col_idx]
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.line_spacing = 1.05
        _run(cp, k + ": ", size=8.9, bold=True, color=ACC2)
        _run(cp, v, size=8.9, color=BODY)
    i += 2

# ---------------- معلومات إضافية ----------------
section_heading("معلومات إضافية", "ADDITIONAL")
rows = [
    ("الجاهزية", "متاح للعمل فورًا بعد استكمال إجراءات التوظيف"),
    ("السكن", "قريب من غرب الرياض — سهولة الوصول للدوام"),
    ("الحالة النظامية", "تأشيرة زيارة عائلية — جاهز لاستكمال الإجراءات"),
    ("المصادر", "لينكدإن والموقع الشخصي أعلاه للاطلاع على الأعمال"),
]
i = 0
while i < len(rows):
    pair = rows[i:i + 2]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    light_bottom(tbl)
    widths = [Cm(8.3), Cm(8.3)]
    for col_idx, (k, v) in enumerate(pair):
        cell = tbl.cell(0, col_idx)
        cell.width = widths[col_idx]
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.line_spacing = 1.05
        _run(cp, k + ": ", size=8.9, bold=True, color=ACC2)
        _run(cp, v, size=8.9, color=BODY)
    i += 2

# ---------------- تذييل ----------------
p = para(before=10)
border_bottom(p, color="E2E8F1", size=4)
_run(p, "", size=1)
p = para(before=2)
_run(p, "سيرة ذاتية — محمد البخيتي  •  موظف مكتب خدمات عامة  •  الرياض",
     size=7.5, color=RGBColor(0x9A, 0xAE, 0xC0))
_run(p, "     +966 56 674 3293", size=8, bold=True, color=ACCENT)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cv-general-services.docx")
doc.save(out)
print("saved:", out)
